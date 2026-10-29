"""Report and audit export files in .xlsx (openpyxl) and PDF (reportlab) (DL-072, owner change
of 2026-10-21 from CSV to .xlsx and PDF; UI-021, UI-087, UI-089).

A file holds one table with the lines every report shows above it: what it counts, the scope and
period, when it was made (Property time) and that amounts are USD with riel as information only
(UI-089, UI-090). A Customer appears only by name, with the Unit (DL-072 answer 6): no file has
an ID number, an address or a document. Amounts are written from integer USD cents. The PDF uses
the built-in Helvetica font, which has no Khmer letters, so Khmer text is readable in the .xlsx
file only.
"""

from __future__ import annotations

import hashlib
import io
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Final, Literal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FileType = Literal["xlsx", "pdf"]
MEDIA_TYPES: Final[Mapping[FileType, str]] = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
Kind = Literal["text", "date", "datetime", "money", "number", "quantity", "bool"]


@dataclass(frozen=True, slots=True)
class Column:
    key: str
    label: str
    kind: Kind = "text"


@dataclass(frozen=True, slots=True)
class ExportFile:
    file_name: str
    media_type: str
    content: bytes
    sha256: str
    row_count: int


C = Column
UNIT = C("unit_number", "Parking")
CUSTOMER = C("customer_name", "Customer")
REGISTER_COLUMNS: Final[Mapping[str, tuple[Column, ...]]] = {
    "invoice_register": (
        C("invoice_number", "Invoice"),
        CUSTOMER,
        UNIT,
        C("primary_category", "Type"),
        C("invoice_date", "Issued", "date"),
        C("due_date", "Due", "date"),
        C("total_usd_cents", "Amount", "money"),
        C("paid_usd_cents", "Paid", "money"),
        C("other_settled_usd_cents", "Credit or Deposit used", "money"),
        C("owed_usd_cents", "Owed", "money"),
        C("overdue_days", "Days overdue", "number"),
        C("status", "Status"),
    ),
    "payments": (
        C("receipt_number", "Receipt"),
        C("payment_date", "Received", "date"),
        CUSTOMER,
        UNIT,
        C("method_name", "Method"),
        C("amount_usd_cents", "Amount", "money"),
        C("received_currency", "Currency received"),
        C("received_amount_minor", "Amount received (riel or cents)", "number"),
        C("rate_khr_per_usd", "Rate (riel per USD)", "number"),
        C("paid_for", "Paid for"),
        C("not_allocated_usd_cents", "Not allocated", "money"),
        C("reversed_on", "Reversed on", "date"),
    ),
    "utility_consumption": (
        UNIT,
        C("utility_type", "Utility"),
        C("period_start", "From", "date"),
        C("period_end", "To", "date"),
        C("quantity", "Used", "quantity"),
        C("is_estimate", "Estimate", "bool"),
        C("amount_usd_cents", "Amount", "money"),
    ),
    "credit_ledger": (
        CUSTOMER,
        UNIT,
        C("receipt_number", "From Receipt"),
        C("recognized_on", "Recorded", "date"),
        C("recognized_usd_cents", "Credit", "money"),
        C("applied_usd_cents", "Used", "money"),
        C("refunded_usd_cents", "Refunded", "money"),
        C("reversed_usd_cents", "Reversed", "money"),
        C("available_usd_cents", "Left", "money"),
    ),
    "deposits": (
        UNIT,
        CUSTOMER,
        C("deposit_type", "Deposit"),
        C("required_usd_cents", "Required", "money"),
        C("received_usd_cents", "Received", "money"),
        C("applied_usd_cents", "Used for debts", "money"),
        C("kept_usd_cents", "Kept", "money"),
        C("refunded_usd_cents", "Refunded", "money"),
        C("held_usd_cents", "Held now", "money"),
        C("record_status", "Status"),
    ),
    "final_settlements": (
        UNIT,
        CUSTOMER,
        C("started_on", "Started", "date"),
        C("move_out_date", "Move-out", "date"),
        C("closed_on", "Closed", "date"),
        C("deposits_applied_usd_cents", "Deposits used", "money"),
        C("credit_applied_usd_cents", "Credit used", "money"),
        C("deposits_kept_usd_cents", "Deposits kept", "money"),
        C("refunded_usd_cents", "Refunded", "money"),
        C("balance_due_usd_cents", "Balance due", "money"),
        C("status", "Status"),
    ),
    "unit_occupancy": (
        UNIT,
        C("unit_type", "Type"),
        C("unit_status", "Status"),
        CUSTOMER,
        C("monthly_rent_usd_cents", "Rent", "money"),
        C("until_date", "Until", "date"),
    ),
    "lease_dates": (
        UNIT,
        CUSTOMER,
        C("start_date", "Start", "date"),
        C("end_date", "End", "date"),
        C("monthly_rent_usd_cents", "Rent", "money"),
        C("days_to_end", "Days to End Date", "number"),
        C("status", "Status"),
    ),
    "audit_activity": (C("day", "Day", "date"), C("count", "Events", "number")),
}
AUDIT_COLUMNS: Final = (
    C("occurred_at", "When", "datetime"),
    C("actor_name", "Who"),
    C("action", "Action"),
    C("entity_type", "Record type"),
    C("entity_id", "Record"),
    UNIT,
    CUSTOMER,
    C("reason", "Reason"),
    C("outcome", "Outcome"),
)
MONTHS: Final = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


def totals_columns(quantities: bool) -> tuple[Column, ...]:
    kind: Kind = "quantity" if quantities else "money"
    first = (C("year", "Year", "number"),) + ((C("what", "What"),) if quantities else ())
    return (
        *first,
        *(C(f"m{number}", name, kind) for number, name in enumerate(MONTHS, start=1)),
        C("total", "Total", kind),
    )


def _plain(value: object, kind: Kind) -> object:
    """The cell value: dollars for money, a number for a quantity, text otherwise."""
    if value is None:
        return None
    if kind == "money":
        return Decimal(int(str(value))) / 100
    if kind == "quantity":
        return Decimal(str(value))
    if kind == "number":
        return int(str(value))
    if kind == "bool":
        return "Yes" if value else "No"
    if kind == "date":
        return value if isinstance(value, date) else date.fromisoformat(str(value)[:10])
    if kind == "datetime":
        return value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
    return str(value)


def _shown(value: object, kind: Kind) -> str:
    """The text a PDF cell shows: dd/mm/yyyy dates and $1,234.56 amounts (DESIGN.md)."""
    cell = _plain(value, kind)
    if cell is None:
        return ""
    if kind == "money":
        amount = Decimal(str(cell))
        return f"-${-amount:,.2f}" if amount < 0 else f"${amount:,.2f}"
    if isinstance(cell, datetime):
        return cell.strftime("%d/%m/%Y %H:%M")
    if isinstance(cell, date):
        return cell.strftime("%d/%m/%Y")
    if isinstance(cell, Decimal):
        return format(cell.normalize(), "f")
    return str(cell)


def _append(sheet: Worksheet, values: Sequence[object]) -> None:
    """Append one row with every text kept as plain text.

    A Customer name or a reason starting with "=" must never become a live formula in the
    spreadsheet, and the control characters a spreadsheet cannot hold are removed."""
    sheet.append(
        [
            ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value
            for value in values
        ]
    )
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.data_type != "s":
            cell.data_type = "s"
            cell.quotePrefix = True


def _xlsx(
    title: str,
    lines: Sequence[str],
    columns: Sequence[Column],
    rows: Sequence[Mapping[str, object]],
) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Report"
    _append(sheet, [title])
    sheet["A1"].font = Font(bold=True, size=14)
    for line in lines:
        _append(sheet, [line])
    sheet.append([])
    _append(sheet, [column.label for column in columns])
    header = sheet.max_row
    for cell in sheet[header]:
        cell.font = Font(bold=True)
    for row in rows:
        _append(sheet, [_plain(row.get(column.key), column.kind) for column in columns])
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        width = 14 if column.kind != "text" else 22
        sheet.column_dimensions[letter].width = width
        style = {
            "money": "$#,##0.00;-$#,##0.00",
            "date": "dd/mm/yyyy",
            "datetime": "dd/mm/yyyy hh:mm",
        }.get(column.kind)
        if style:
            for (cell,) in sheet.iter_rows(min_row=header + 1, min_col=index, max_col=index):
                cell.number_format = style
    sheet.freeze_panes = f"A{header + 1}"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _pdf(
    title: str,
    lines: Sequence[str],
    columns: Sequence[Column],
    rows: Sequence[Mapping[str, object]],
) -> bytes:
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=title,
        author="CHETANA Parking",
        invariant=True,
    )
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=7, leading=8.5)
    head_style = cell_style.clone("head", fontName="Helvetica-Bold")
    story: list[Flowable] = [Paragraph(_escape(title), styles["Title"])]
    story.extend(Paragraph(_escape(line), styles["BodyText"]) for line in lines)
    story.append(Spacer(1, 4 * mm))
    data = [[Paragraph(_escape(column.label), head_style) for column in columns]]
    data.extend(
        [
            Paragraph(_escape(_shown(row.get(column.key), column.kind)), cell_style)
            for column in columns
        ]
        for row in rows
    )
    table = Table(data, repeatRows=1)
    style: list[tuple[Any, ...]] = [
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EFE9")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    style.extend(
        ("ALIGN", (index, 0), (index, -1), "RIGHT")
        for index, column in enumerate(columns)
        if column.kind in {"money", "number", "quantity"}
    )
    table.setStyle(TableStyle(style))
    story.append(table)
    document.build(story)
    return output.getvalue()


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_file(
    file_type: FileType,
    base_name: str,
    title: str,
    lines: Sequence[str],
    columns: Sequence[Column],
    rows: Sequence[Mapping[str, object]],
) -> ExportFile:
    """One export file with its SHA-256 fingerprint (DL-072 answer 12)."""
    content = (_xlsx if file_type == "xlsx" else _pdf)(title, lines, columns, rows)
    return ExportFile(
        file_name=f"{base_name}.{file_type}",
        media_type=MEDIA_TYPES[file_type],
        content=content,
        sha256=hashlib.sha256(content).hexdigest(),
        row_count=len(rows),
    )
