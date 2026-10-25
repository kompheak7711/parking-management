"""V1-057 step 4: report and audit export files (DL-072 answers 3 to 6, 11, 12 and 14).

A report file needs `data.export` and the report's own read rights; the audit file needs
`audit.export`. A small file downloads at once and is not kept; a larger one is kept as an
export job. Every export writes an audit event with the file's SHA-256 fingerprint. A file shows
a Customer only by name and Unit, and Staff-level Utility files have no amounts.
"""

from __future__ import annotations

import hashlib
import io
import uuid
from collections.abc import Sequence
from datetime import date, datetime, timedelta
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from property_rental.app import create_app
from property_rental.audit import _freeze_audit_value
from property_rental.identity import AuthenticatedActor, CapabilityGrant
from property_rental.report_exports import AuditExportRequest, paid_for
from property_rental.report_files import ExportFile
from property_rental.settings import Settings

PROPERTY = uuid.uuid4()
U1 = uuid.uuid4()
UNITS: list[dict[str, object]] = [
    {"id": U1, "unit_number": "U101", "floor": "1", "unit_type": "STUDIO", "status": "OCCUPIED"}
]
BROWSER = {
    "Origin": "https://localhost",
    "Cookie": "__Host-prm_csrf=csrf-test-token",
    "X-CSRF-Token": "csrf-test-token",
}
GENERATED = datetime(2026, 10, 25, 18, 30)


def invoice_row(number: int) -> dict[str, object]:
    return {
        "invoice_id": uuid.uuid4(),
        "invoice_number": f"INV-{number:05d}",
        "status": "UNPAID",
        "customer_id": uuid.uuid4(),
        "customer_name": "Sok Dara",
        "lease_id": uuid.uuid4(),
        "unit_id": U1,
        "primary_category": "RENT",
        "rent_billing_kind": "RECURRING",
        "invoice_date": date(2026, 9, 25),
        "due_date": date(2026, 10, 5),
        "total_usd_cents": 30_000,
        "paid_usd_cents": 10_000,
        "other_settled_usd_cents": 0,
        "owed_usd_cents": 20_000,
        "overdue_days": 24,
    }


class FakeReportStore:
    def __init__(self, invoices: int = 2) -> None:
        self.invoices = invoices

    def home_page(
        self, actor_id: uuid.UUID, property_id: uuid.UUID, sections: Sequence[str]
    ) -> dict[str, object]:
        raise AssertionError("not used")

    def rows(
        self,
        read: str,
        actor_id: uuid.UUID,
        property_id: uuid.UUID,
        extra: Sequence[object] = (),
    ) -> list[dict[str, object]]:
        if read == "report_invoice_register":
            return [invoice_row(number) for number in range(self.invoices)]
        if read == "report_utility_facts":
            return [
                {
                    "period_id": uuid.uuid4(),
                    "unit_id": U1,
                    "month": date(2026, 10, 1),
                    "utility_type": "WATER",
                    "period_start": date(2026, 9, 25),
                    "period_end": date(2026, 10, 24),
                    "quantity": "6.5",
                    "is_estimate": False,
                    "amount_usd_cents": 650,
                    "billed": True,
                }
            ]
        return []

    def units(self, property_id: uuid.UUID) -> list[dict[str, object]]:
        return list(UNITS)

    def utility_rates(self, property_id: uuid.UUID) -> list[dict[str, object]]:
        return []


class FakeExportStore:
    def __init__(self) -> None:
        self.records: list[tuple[str, str, dict[str, object]]] = []
        self.kept: dict[uuid.UUID, dict[str, object]] = {}
        self.audit_requests: list[AuditExportRequest] = []

    def generated_at(self, property_id: uuid.UUID) -> datetime:
        return GENERATED

    def record(
        self,
        actor_id: uuid.UUID,
        property_id: uuid.UUID,
        kind: str,
        capability: str,
        details: dict[str, object],
    ) -> None:
        self.records.append((kind, capability, details))

    def save(
        self,
        actor_id: uuid.UUID,
        property_id: uuid.UUID,
        export_id: uuid.UUID,
        report: str,
        file_type: str,
        file: ExportFile,
        kind: str,
        capability: str,
        details: dict[str, object],
    ) -> datetime:
        self.records.append((kind, capability, details))
        self.kept[export_id] = {
            "actor_id": actor_id,
            "report": report,
            "file_type": file_type,
            "file_name": file.file_name,
            "media_type": file.media_type,
            "content": file.content,
            "sha256": file.sha256,
            "row_count": file.row_count,
        }
        return GENERATED + timedelta(hours=24)

    def take(
        self, actor_id: uuid.UUID, property_id: uuid.UUID, export_id: uuid.UUID
    ) -> dict[str, object] | None:
        kept = self.kept.get(export_id)
        return kept if kept is not None and kept["actor_id"] == actor_id else None

    def audit_rows(
        self, property_id: uuid.UUID, request: AuditExportRequest
    ) -> list[dict[str, object]]:
        self.audit_requests.append(request)
        return [
            {
                "occurred_at": datetime(2026, 10, 25, 9, 0),
                "actor_name": "Chan",
                "action": "invoice.issued",
                "entity_type": "invoice",
                "entity_id": uuid.uuid4(),
                "unit_number": "U101",
                "customer_name": "Sok Dara",
                "reason": "Billing run",
                "outcome": "SUCCEEDED",
            }
        ]


def person(*capabilities: str) -> AuthenticatedActor:
    return AuthenticatedActor(
        actor_id=str(uuid.uuid4()),
        capability_grants=frozenset(
            CapabilityGrant(capability, frozenset({str(PROPERTY)})) for capability in capabilities
        ),
    )


def client(
    actor: AuthenticatedActor | None,
    reports: FakeReportStore | None = None,
    exports: FakeExportStore | None = None,
) -> TestClient:
    return TestClient(
        create_app(
            Settings(database_url=None),
            actor_resolver=lambda _: actor,
            report_store=reports or FakeReportStore(),
            export_store=exports or FakeExportStore(),
        )
    )


EXPORTS = f"/api/v1/properties/{PROPERTY}/exports"
AUDIT = f"/api/v1/properties/{PROPERTY}/audit-exports"
MANAGER = ("reports.read", "data.export", "utility.read")


def sheet_rows(content: bytes) -> list[tuple[Any, ...]]:
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    rows = []
    for row in sheet.iter_rows(values_only=True):
        cells = list(row)
        while cells and cells[-1] is None:
            cells.pop()
        rows.append(tuple(cells))
    return rows


def test_a_report_file_needs_data_export_and_the_report_right() -> None:
    exports = FakeExportStore()
    body = {"report": "invoice_register", "file_type": "xlsx"}
    assert (
        client(None, exports=exports).post(EXPORTS, json=body, headers=BROWSER).status_code == 401
    )
    reader = client(person("reports.read"), exports=exports)
    assert reader.post(EXPORTS, json=body, headers=BROWSER).status_code == 403
    staff_export = client(person("data.export", "invoices.read"), exports=exports)
    assert staff_export.post(EXPORTS, json=body, headers=BROWSER).status_code == 403
    manager = client(person(*MANAGER), exports=exports)
    assert manager.post(EXPORTS, json=body).status_code == 403
    assert exports.records == []


def test_a_small_xlsx_file_downloads_at_once_and_its_event_keeps_its_fingerprint() -> None:
    exports = FakeExportStore()
    manager = client(person(*MANAGER), exports=exports)
    answer = manager.post(
        EXPORTS,
        json={
            "report": "invoice_register",
            "file_type": "xlsx",
            "floor": "1",
            "from": "2026-09-01",
        },
        headers=BROWSER,
    )
    assert answer.status_code == 200
    assert answer.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="invoice_register_20261025_1830.xlsx"' in answer.headers["content-disposition"]
    digest = hashlib.sha256(answer.content).hexdigest()
    assert answer.headers["x-export-sha256"] == digest
    rows = sheet_rows(answer.content)
    assert rows[0] == ("Invoice register and balances",)
    assert ("Scope: Floor 1",) in rows and ("Period: From 01/09/2026",) in rows
    header = rows.index(next(row for row in rows if row and row[0] == "Invoice"))
    assert rows[header][:3] == ("Invoice", "Customer", "Unit")
    assert rows[header + 1][1:3] == ("Sok Dara", "U101")
    assert float(rows[header + 1][6]) == 300.0
    [(kind, capability, details)] = exports.records
    assert (kind, capability) == ("report_export", "data.export")
    assert details["file_sha256"] == digest and details["delivery_kind"] == "download"
    assert details["row_count"] == 2 and details["report_kind"] == "invoice_register"
    assert details["filters"] == {"floor_label": "1", "from_date": "2026-09-01"}
    assert exports.kept == {}


def test_a_pdf_file_and_a_large_export_job_kept_for_its_maker() -> None:
    exports = FakeExportStore()
    maker = person(*MANAGER)
    manager = client(maker, FakeReportStore(invoices=1001), exports)
    pdf = manager.post(
        EXPORTS, json={"report": "unit_occupancy", "file_type": "pdf"}, headers=BROWSER
    )
    assert pdf.status_code == 403  # Unit occupancy also needs its read rights.
    job = manager.post(
        EXPORTS, json={"report": "invoice_register", "file_type": "pdf"}, headers=BROWSER
    )
    assert job.status_code == 201
    data = job.json()["data"]
    assert data["row_count"] == 1001 and data["expires_at"].startswith("2026-10-26")
    assert exports.records[-1][2]["delivery_kind"] == "job"
    # Every detail passes the audit value rules, so the History page can show the event.
    for _, _, recorded in exports.records:
        _freeze_audit_value(recorded, "resulting_state")
    kept = manager.get(f"{EXPORTS}/{data['export_id']}")
    assert kept.status_code == 200 and kept.content.startswith(b"%PDF")
    assert hashlib.sha256(kept.content).hexdigest() == data["sha256"]
    other = client(person(*MANAGER), FakeReportStore(), exports)
    assert other.get(f"{EXPORTS}/{data['export_id']}").status_code == 404
    assert (
        client(person("reports.read"), exports=exports)
        .get(f"{EXPORTS}/{data['export_id']}")
        .status_code
        == 403
    )


def test_a_utility_file_has_amounts_only_with_the_report_right() -> None:
    staff = client(person("utility.read", "data.export"))
    answer = staff.post(
        EXPORTS, json={"report": "utility_consumption", "file_type": "xlsx"}, headers=BROWSER
    )
    assert answer.status_code == 200
    rows = sheet_rows(answer.content)
    header = next(row for row in rows if row and row[0] == "Unit")
    assert "Amount" not in header and "Used" in header
    assert ("Amounts in USD; riel shown as information only.",) in rows
    manager = client(person(*MANAGER))
    totals = manager.post(
        EXPORTS, json={"report": "total_water", "file_type": "xlsx"}, headers=BROWSER
    )
    assert totals.status_code == 200
    total_rows = sheet_rows(totals.content)
    assert ("Year", "What", "Jan") == next(row for row in total_rows if row and row[0] == "Year")[
        :3
    ]


def test_the_export_body_is_bounded() -> None:
    manager = client(person(*MANAGER))
    for body in (
        {"report": "invoice_register", "file_type": "csv"},
        {"report": "customer_list", "file_type": "xlsx"},
        {"report": "invoice_register", "file_type": "xlsx", "unit_id": str(U1)},
        {"report": "invoice_register", "file_type": "xlsx", "id_number": "1"},
        {
            "report": "invoice_register",
            "file_type": "xlsx",
            "from": "2026-05-01",
            "to": "2026-04-01",
        },
    ):
        assert manager.post(EXPORTS, json=body, headers=BROWSER).status_code == 422, body


def test_the_audit_file_needs_audit_export_and_keeps_its_fingerprint() -> None:
    exports = FakeExportStore()
    body = {"file_type": "xlsx", "from": "2026-10-01", "action": "invoice.issued"}
    viewer = client(person("audit.view", "data.export"), exports=exports)
    assert viewer.post(AUDIT, json=body, headers=BROWSER).status_code == 403
    owner = client(person("audit.export"), exports=exports)
    answer = owner.post(AUDIT, json=body, headers=BROWSER)
    assert answer.status_code == 200
    [(kind, capability, details)] = exports.records
    assert (kind, capability) == ("audit_export", "audit.export")
    assert details["file_sha256"] == hashlib.sha256(answer.content).hexdigest()
    assert exports.audit_requests[0].action == "invoice.issued"
    rows = sheet_rows(answer.content)
    assert rows[0] == ("Audit history",)
    header = next(row for row in rows if row and row[0] == "When")
    assert header == (
        "When",
        "Who",
        "Action",
        "Record type",
        "Record",
        "Unit",
        "Customer",
        "Reason",
        "Outcome",
    )
    assert not any("riel" in str(row[0]) for row in rows if row)


def test_paid_for_names_each_line_with_its_amount() -> None:
    parts = [
        {"kind": "INVOICES", "category": "RENT", "amount_usd_cents": 30_000},
        {"kind": "NOT_ALLOCATED", "category": None, "amount_usd_cents": 1_050},
    ]
    assert paid_for(parts) == "Invoices Rent $300.00; Not allocated $10.50"
    assert paid_for(None) == ""
