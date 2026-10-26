"""V1-058 R4c: reading a meter-readings spreadsheet and making the blank one."""

from __future__ import annotations

import base64
import io
import zipfile
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from property_rental.utility_sheet import (
    MAX_ROWS,
    SheetError,
    SheetRow,
    blank_sheet,
    read_sheet,
)
from test_utility_api import InMemoryUtilityFlow, browser_headers, client_for


def xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_an_xlsx_file_gives_one_row_per_unit_with_the_text_people_see() -> None:
    content = xlsx(
        [
            [None, None, None],
            ["Water", "Unit", "Electricity"],
            [12, "U101", 2430],
            [None, "U102", 3142.0],
            [None, None, None],
            [7.5, " U103 ", None],
        ]
    )
    assert read_sheet("readings-24-11-2026.XLSX", content) == [
        SheetRow(row=3, unit="U101", electricity="2430", water="12"),
        SheetRow(row=4, unit="U102", electricity="3142", water=""),
        SheetRow(row=6, unit="U103", electricity="", water="7.5"),
    ]


def test_a_csv_file_with_a_byte_order_mark_reads_the_same_way() -> None:
    content = "﻿unit,ELECTRICITY,water\r\nU101,2430,12\r\n,,\r\nU102,,9\r\n".encode()
    assert read_sheet("r.csv", content) == [
        SheetRow(row=2, unit="U101", electricity="2430", water="12"),
        SheetRow(row=4, unit="U102", electricity="", water="9"),
    ]


@pytest.mark.parametrize(
    ("name", "content", "problem"),
    [
        ("r.xls", b"x", "FILE_TYPE"),
        ("r.xlsx", b"not a zip", "UNREADABLE"),
        ("r.csv", b"\xff\xfe\x00", "UNREADABLE"),
        ("r.csv", b"Unit,Electricity\nU101,1\n", "COLUMNS"),
        ("r.csv", b"Unit,Electricity,Water\n", "EMPTY"),
        ("r.csv", b"", "EMPTY"),
        (
            "r.csv",
            ("Unit,Electricity,Water\n" + "U1,1,1\n" * (MAX_ROWS + 1)).encode(),
            "TOO_MANY_ROWS",
        ),
    ],
)
def test_a_file_that_cannot_be_used_says_why(name: str, content: bytes, problem: str) -> None:
    with pytest.raises(SheetError) as raised:
        read_sheet(name, content)
    assert raised.value.problem == problem


def test_the_blank_file_lists_every_unit_under_the_three_columns() -> None:
    workbook = load_workbook(io.BytesIO(blank_sheet(["U101", "U102"])))
    sheet = workbook.worksheets[0]
    assert [list(row) for row in sheet.iter_rows(values_only=True)] == [
        ["Unit", "Electricity", "Water"],
        ["U101", None, None],
        ["U102", None, None],
    ]
    assert read_sheet("blank.xlsx", blank_sheet(["U101"])) == [
        SheetRow(row=2, unit="U101", electricity="", water="")
    ]


def test_a_unit_that_looks_like_a_formula_stays_text_in_the_blank_file() -> None:
    sheet = load_workbook(io.BytesIO(blank_sheet(["=1+1"]))).worksheets[0]
    assert sheet["A2"].data_type == "s"
    assert sheet["A2"].value == "=1+1"


def test_empty_rows_are_not_counted_and_a_semicolon_csv_reads_the_same_way() -> None:
    padded = "Unit,Electricity,Water\nU101,1,2\n" + ",,\n" * (MAX_ROWS + 100)
    assert read_sheet("r.csv", padded.encode()) == [
        SheetRow(row=2, unit="U101", electricity="1", water="2")
    ]
    rows: list[list[object]] = [["Unit", "Electricity", "Water"], ["U101", 1, 2]]
    formatted = xlsx(rows + [[None]] * 1000)
    assert len(read_sheet("r.xlsx", formatted)) == 1
    assert read_sheet("r.csv", b"Unit;Electricity;Water\nU101;1300;12\n") == [
        SheetRow(row=2, unit="U101", electricity="1300", water="12")
    ]


def _replace_part(content: bytes, name: str, data: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(content))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item.filename, data if item.filename == name else source.read(item))
    return buffer.getvalue()


def test_a_broken_or_swollen_file_is_unreadable_not_a_server_error() -> None:
    good = xlsx([["Unit", "Electricity", "Water"], ["U101", 1, 2]])
    cases = [
        ("r.csv", ("Unit,Electricity,Water\nU101," + "9" * 200_000 + ",1\n").encode()),
        ("r.xlsx", _replace_part(good, "xl/worksheets/sheet1.xml", b"<worksheet><broken")),
        ("r.xlsx", _replace_part(good, "xl/workbook.xml", b"<workbook")),
        # A part that unpacks to far more than any readings sheet needs.
        ("r.xlsx", _replace_part(good, "xl/theme/theme1.xml", b" " * (11 * 1024 * 1024))),
    ]
    for name, content in cases:
        assert len(content) < 2 * 1024 * 1024
        with pytest.raises(SheetError) as raised:
            read_sheet(name, content)
        assert raised.value.problem == "UNREADABLE"


def test_the_sheet_route_names_the_problem_of_a_file() -> None:
    flow = InMemoryUtilityFlow(uuid4())
    recorder = client_for(flow, "utility.read", "utility_readings.record")
    response = recorder.post(
        f"/api/v1/properties/{flow.property_id}/utility-readings/sheet",
        json={
            "file_name": "r.csv",
            "content_base64": base64.b64encode(b"Unit,Meter\nU101,1\n").decode(),
        },
        headers=browser_headers(),
    )
    assert response.status_code == 200
    assert response.json()["data"] == {"file_name": "r.csv", "problem": "COLUMNS", "rows": []}


def test_the_sheet_routes_need_their_rights_and_the_browser_checks() -> None:
    flow = InMemoryUtilityFlow(uuid4())
    sheet_path = f"/api/v1/properties/{flow.property_id}/utility-readings/sheet"
    blank_path = f"/api/v1/properties/{flow.property_id}/utility-readings/blank-sheet"
    body = {
        "file_name": "r.csv",
        "content_base64": base64.b64encode(b"Unit,Electricity,Water\nU101,5,6\n").decode(),
    }
    reader = client_for(flow, "utility.read")
    assert reader.post(sheet_path, json=body, headers=browser_headers()).status_code == 403
    recorder = client_for(flow, "utility.read", "utility_readings.record")
    assert recorder.post(sheet_path, json=body).status_code == 403
    response = recorder.post(sheet_path, json=body, headers=browser_headers())
    assert response.status_code == 200
    assert response.json()["data"] == {
        "file_name": "r.csv",
        "problem": None,
        "rows": [{"row": 2, "unit": "U101", "electricity": "5", "water": "6"}],
    }
    broken = recorder.post(
        sheet_path,
        json={"file_name": "r.csv", "content_base64": "%%%"},
        headers=browser_headers(),
    )
    assert broken.json()["data"]["problem"] == "UNREADABLE"
    assert (
        recorder.post(
            f"/api/v1/properties/{uuid4()}/utility-readings/sheet",
            json=body,
            headers=browser_headers(),
        ).status_code
        == 403
    )

    assert (
        client_for(flow)
        .post(blank_path, json={"units": ["U101"]}, headers=browser_headers())
        .status_code
        == 403
    )
    assert reader.post(blank_path, json={"units": ["U101"]}).status_code == 403
    blank = reader.post(blank_path, json={"units": ["U101"]}, headers=browser_headers())
    assert blank.status_code == 200
    assert blank.headers["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert read_sheet("blank.xlsx", blank.content)[0].unit == "U101"
    assert reader.post(blank_path, json={"units": []}, headers=browser_headers()).status_code == 422
